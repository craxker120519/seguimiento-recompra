"""Pestaña Resultados: estadisticas por caso, facturas nuevas y datos de demostracion."""

import io
import json
import os

import openpyxl
import pytest

from app import negocio, server
from app.herramientas import datos_demo
from app.negocio import resultados as neg_resultados
from app.web import comun
from tests.conftest import iso


def test_cada_caso_cuenta_con_su_resultado_mas_reciente(seg):
    seg("PLASTICOS", "no_contesto", 5)
    seg("PLASTICOS", "no_contesto", 4)
    seg("PLASTICOS", "vendio_de_nuevo", 3)
    seg("TALLER", "no_contesto", 6)
    seg("TALLER", "cambio_proveedor", 5, nota="mas barato")
    seg("OTRO", "contactado_esperar", 2, recordar="2026-12-01")
    seg("CERRO", "no_aplica", 1, nota="cerro")
    r = negocio.resumen_seguimientos(90)
    assert r["total"] == 4
    assert r["totales"] == {
        "contactado_esperar": 1,
        "vendio_de_nuevo": 1,
        "cambio_proveedor": 1,
        "no_contesto": 0,
        "no_aplica": 1,
    }
    casos = {c["clave_cliente"]: c for c in r["casos"]}
    assert casos["PLASTICOS"]["intentos"] == 2 and casos["PLASTICOS"]["seguimientos"] == 3
    assert r["tasa_recuperacion"] == 50  # 1 de 2 casos resueltos; "no aplica" no cuenta
    assert [c["fecha"] for c in r["casos"]] == sorted((c["fecha"] for c in r["casos"]), reverse=True)


def test_periodo_y_base_vacia(seg):
    assert negocio.resumen_seguimientos(30)["tasa_recuperacion"] is None
    seg("VIEJO", "vendio_de_nuevo", 60)
    assert negocio.resumen_seguimientos(30)["total"] == 0
    assert negocio.resumen_seguimientos(90)["total"] == 1
    assert negocio.resumen_seguimientos(None)["total"] == 1


def test_una_factura_nueva_recupera_perdidos_pendientes_y_descartados(seg):
    seg("PERDIDO", "cambio_proveedor", 10)
    seg("ESPERA", "contactado_esperar", 10, recordar="2026-12-01")
    seg("DESCARTADO", "no_aplica", 10, nota="x")
    seg("MUDO", "no_contesto", 10)
    seg("SIN FACTURA", "cambio_proveedor", 10)
    d = negocio.resumen_seguimientos(90)
    compras = {
        ("PERDIDO", "A"): iso(2),
        ("ESPERA", "A"): iso(2),
        ("DESCARTADO", "A"): iso(2),
        ("MUDO", "A"): iso(30),
    }  # factura anterior: no cuenta
    negocio.reclasificar_por_compras(d, compras)
    casos = {c["clave_cliente"]: c for c in d["casos"]}
    for cliente, antes in (
        ("PERDIDO", "cambio_proveedor"),
        ("ESPERA", "contactado_esperar"),
        ("DESCARTADO", "no_aplica"),
    ):
        assert (
            casos[cliente]["resultado"] == "vendio_de_nuevo"
            and casos[cliente]["antes"] == antes
            and casos[cliente]["por_factura"]
        )
    assert casos["MUDO"]["resultado"] == "no_contesto"
    assert casos["SIN FACTURA"]["resultado"] == "cambio_proveedor"
    assert d["totales"]["vendio_de_nuevo"] == 3 and d["tasa_recuperacion"] == 75


def test_ventas_sin_factura_se_marcan_en_resultados(seg):
    seg("ESPERANDO", "vendio_de_nuevo", 5)
    seg("VENCIDA", "vendio_de_nuevo", 20)
    seg("FACTURADA", "vendio_de_nuevo", 5)
    d = negocio.resumen_seguimientos(90)
    negocio.reclasificar_por_compras(d, {("FACTURADA", "A"): iso(0)})
    c = {x["clave_cliente"]: x for x in d["casos"]}
    assert c["ESPERANDO"]["sin_factura"] and not c["ESPERANDO"]["factura_vencida"]
    assert c["VENCIDA"]["sin_factura"] and c["VENCIDA"]["factura_vencida"]
    assert not c["FACTURADA"].get("sin_factura")


def test_ventas_registradas_solo_cuenta_las_que_terminaron_en_volvio_a_comprar(seg):
    seg("A", "vendio_de_nuevo", 10)
    seg("B", "vendio_de_nuevo", 10)
    seg("B", "cambio_proveedor", 5)
    seg("DEMO - X", "vendio_de_nuevo", 10)  # los clientes de la demo nunca avisan
    assert [v["clave_cliente"] for v in negocio.ventas_registradas()] == ["A"]


# ---- endpoints ----


@pytest.fixture
def cliente_web(monkeypatch):
    monkeypatch.setattr(comun, "ultimas_compras", lambda codigos: {})
    monkeypatch.setattr(comun, "codigos_filtro", lambda: ["A"])
    monkeypatch.setattr(comun, "catalogo_filtros", lambda codigos: [{"codigo": "A", "nombre": "FILTRO A"}])
    c = server.app.test_client()
    with c.session_transaction() as s:
        s.update(usuario_id=1, usuario_rol="admin", usuario_nombre="Ana", usuario_username="ana")
    return c


def test_endpoint_y_excel_de_resultados(cliente_web, seg, monkeypatch):
    seg("KIPP", "vendio_de_nuevo", 5)
    seg("MOLINOS", "cambio_proveedor", 5, nota="mas barato")
    monkeypatch.setattr(comun, "ultimas_compras", lambda codigos: {("MOLINOS", "A"): iso(1)})
    j = cliente_web.get("/api/resultados?dias=90").get_json()
    assert j["total"] == 2 and j["totales"]["vendio_de_nuevo"] == 2
    assert {c["clave_cliente"]: c.get("por_factura", False) for c in j["casos"]} == {"KIPP": False, "MOLINOS": True}
    assert all(c["nombre_producto"] == "FILTRO A" for c in j["casos"])
    wb = openpyxl.load_workbook(io.BytesIO(cliente_web.get("/api/exportar/resultados.xlsx?dias=90").data))
    assert wb.sheetnames == ["Resumen", "Clientes", "Historial"]
    notas = [r[7] for r in wb["Clientes"].iter_rows(min_row=2, values_only=True)]
    assert any(n and n.startswith("Factura nueva detectada (antes: Cambió de proveedor)") for n in notas)
    assert cliente_web.get("/api/resultados?dias=0").status_code == 400
    assert cliente_web.get("/api/exportar/resultados.xlsx?dias=-3").status_code == 400


def test_resultados_sigue_funcionando_si_erp_no_responde(cliente_web, seg, monkeypatch):
    seg("KIPP", "cambio_proveedor", 5)

    def falla(codigos):
        raise RuntimeError("sin conexion")

    monkeypatch.setattr(comun, "ultimas_compras", falla)
    j = cliente_web.get("/api/resultados").get_json()
    assert j["totales"]["cambio_proveedor"] == 1


def test_las_rutas_exigen_login():
    c = server.app.test_client()
    assert c.get("/api/resultados").status_code == 401
    assert c.post("/api/seguimiento", json={}).status_code == 401


# ---- demo ----


def test_la_demo_muestra_toda_la_funcion_y_se_borra(cliente_web, monkeypatch):
    datos_demo.crear()
    assert os.path.exists(neg_resultados.RUTA_COMPRAS_DEMO)
    monkeypatch.setattr(comun, "ultimas_compras", lambda codigos: db_demo_compras())
    j = cliente_web.get("/api/resultados?dias=90").get_json()
    assert all(v > 0 for k, v in j["totales"].items() if k != "no_aplica")
    casos = {c["clave_cliente"].replace("DEMO - ", ""): c for c in j["casos"]}
    assert casos["COMERCIAL DEL PACIFICO"]["factura_vencida"]
    assert casos["DISTRIBUIDORA ANAHUAC"]["sin_factura"] and not casos["DISTRIBUIDORA ANAHUAC"]["factura_vencida"]
    for nombre in ("FERRETERIA CENTRAL", "MOLINOS DEL NORTE", "CLINICA SAN RAFAEL"):
        assert casos[nombre]["por_factura"]
    assert negocio.ventas_registradas() == []  # la demo nunca dispara avisos
    datos_demo.borrar()
    assert not os.path.exists(neg_resultados.RUTA_COMPRAS_DEMO)
    assert negocio.resumen_seguimientos(None)["total"] == 0


def db_demo_compras():
    with open(neg_resultados.RUTA_COMPRAS_DEMO, encoding="utf-8") as f:
        return {tuple(k.split("|", 1)): v for k, v in json.load(f).items()}
