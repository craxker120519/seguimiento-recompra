"""Descargas en Excel: contactar hoy, catalogo, un producto, un cliente y resultados."""

import io
from datetime import date

import openpyxl
from flask import Blueprint, jsonify, request, send_file

from app import negocio
from app.erp import alertas_recontacto, catalogo_filtros, timeline_cliente, timeline_producto
from app.negocio.filtros import codigos_filtro
from app.web.comun import ETIQUETAS_RESULTADO, datos_resultados, dias_de_la_peticion, nombres_producto

bp = Blueprint("exportaciones", __name__)

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_CARACTERES_INVALIDOS_HOJA = set(":\\/?*[]")
_NEGRITA = openpyxl.styles.Font(bold=True)


# ---- ayudas ----


def _nombre_hoja_valido(texto, usados):
    """Excel limita el nombre de una hoja a 31 caracteres, prohibe : \\ / ? * [ ]
    y no permite hojas repetidas; aqui se sanea y se desempata."""
    limpio = "".join(c for c in str(texto) if c not in _CARACTERES_INVALIDOS_HOJA)[:31].strip() or "Hoja"
    candidato = limpio
    i = 2
    while candidato.lower() in usados:
        sufijo = f" ({i})"
        candidato = limpio[: 31 - len(sufijo)] + sufijo
        i += 1
    usados.add(candidato.lower())
    return candidato


def _hipervinculo(celda, hoja_destino):
    """Convierte una celda en un link interno a la celda A1 de otra hoja
    del mismo libro (para saltar del resumen al desglose y viceversa)."""
    celda.hyperlink = f"#'{hoja_destino}'!A1"
    celda.font = openpyxl.styles.Font(color="0563C1", underline="single")


def _anchos(hoja, anchos):
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ancho


def _encabezado_en_negrita(fila):
    for c in fila:
        c.font = _NEGRITA


def _descargar(libro, nombre_archivo):
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name=nombre_archivo, mimetype=TIPO_XLSX)


def _excel_de_una_hoja(hoja_titulo, encabezados, filas, anchos, nombre_archivo):
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = hoja_titulo
    hoja.append(encabezados)
    _encabezado_en_negrita(hoja[1])
    for fila in filas:
        hoja.append(fila)
    _anchos(hoja, anchos)
    return _descargar(libro, nombre_archivo)


def _libro_resumen_y_detalle(elementos, nombre_hoja, encabezados, anchos_resumen, fila_resumen):
    """Hoja "Resumen" con un renglon por elemento (con link a su hoja de detalle), mas una hoja por
    elemento con su linea de tiempo completa de compras (con link de vuelta al resumen)."""
    nombres_usados = {"resumen"}
    hojas = [_nombre_hoja_valido(nombre_hoja(e), nombres_usados) for e in elementos]

    libro = openpyxl.Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    resumen.append(encabezados)
    _encabezado_en_negrita(resumen[1])
    for elemento, hoja in zip(elementos, hojas, strict=True):
        ultima = elemento.get("compras", [{}])[0]
        resumen.append(fila_resumen(elemento, ultima))
        _hipervinculo(resumen.cell(row=resumen.max_row, column=1), hoja)
    _anchos(resumen, anchos_resumen)

    for elemento, nombre in zip(elementos, hojas, strict=True):
        detalle = libro.create_sheet(nombre)
        detalle.append(["← Volver al resumen"])
        _hipervinculo(detalle.cell(row=1, column=1), "Resumen")
        detalle.append(["Fecha", "Unidades", "Precio unitario", "Moneda", "Total", "Próximo contacto"])
        _encabezado_en_negrita(detalle[2])
        for compra in elemento.get("compras", []):
            detalle.append(
                [
                    compra.get("fecha", ""),
                    compra.get("unidades", ""),
                    compra.get("precio_unitario", ""),
                    compra.get("moneda", ""),
                    compra.get("monto", ""),
                    compra.get("proximo_contacto") or "",
                ]
            )
        _anchos(detalle, [14, 10, 14, 8, 12, 16])
    return libro


# ---- rutas ----


@bp.route("/api/exportar/alertas.xlsx")
def api_exportar_alertas():
    try:
        alertas = alertas_recontacto(codigos_filtro(), incluir_esperando_factura=True)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

    def situacion(a):
        if a["estado"] == "esperando_factura":
            return "Esperando factura"
        if a["estado"] == "factura_vencida":
            return "Sin factura tras 15 días: contactar"
        return "Vencido" if a["vencido"] else "Por vencer"

    filas = [
        [
            a["cliente"],
            a["codigo"],
            a["nombre_producto"],
            a["ultima_compra"],
            a["proximo_contacto"],
            "Sí" if a["vencido"] else "No",
            situacion(a),
        ]
        for a in alertas
    ]
    return _excel_de_una_hoja(
        "Contactar hoy",
        ["Cliente", "Código", "Producto", "Última compra", "Próximo contacto", "Vencido", "Situación"],
        filas,
        [32, 12, 40, 14, 16, 9, 34],
        f"contactar-hoy-{date.today().isoformat()}.xlsx",
    )


@bp.route("/api/exportar/catalogo.xlsx")
def api_exportar_catalogo():
    try:
        catalogo = catalogo_filtros(codigos_filtro(), con_vida_util=True)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

    filas = [[f["codigo"], f.get("nombre", ""), f.get("vida_util_meses", "")] for f in catalogo]
    return _excel_de_una_hoja(
        "Catálogo",
        ["Código", "Nombre", "Vida útil (meses)"],
        filas,
        [14, 50, 16],
        f"catalogo-filtros-{date.today().isoformat()}.xlsx",
    )


@bp.route("/api/exportar/producto/<codigo>.xlsx")
def api_exportar_producto(codigo):
    try:
        d = timeline_producto(codigo)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

    libro = _libro_resumen_y_detalle(
        d.get("clientes", []),
        nombre_hoja=lambda cli: cli.get("cliente", "cliente"),
        encabezados=["Cliente", "Unidades totales", "Última compra", "Próximo contacto"],
        anchos_resumen=[32, 16, 14, 16],
        fila_resumen=lambda cli, ultima: [
            cli.get("cliente", ""),
            cli.get("unidades_total", ""),
            ultima.get("fecha", ""),
            ultima.get("proximo_contacto") or "",
        ],
    )
    return _descargar(libro, f"producto-{codigo}-{date.today().isoformat()}.xlsx")


@bp.route("/api/exportar/cliente.xlsx")
def api_exportar_cliente():
    legacy_id = request.args.get("legacy_id", type=int)
    actual_id = request.args.get("actual_id", type=int)
    if legacy_id is None and actual_id is None:
        return jsonify({"error": "falta legacy_id o actual_id"}), 400
    try:
        d = timeline_cliente(codigos_filtro(), legacy_id=legacy_id, actual_id=actual_id)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

    libro = _libro_resumen_y_detalle(
        d.get("productos", []),
        nombre_hoja=lambda p: p.get("codigo", "producto"),
        encabezados=["Código", "Producto", "Unidades totales", "Última compra", "Próximo contacto"],
        anchos_resumen=[14, 40, 16, 14, 16],
        fila_resumen=lambda p, ultima: [
            p.get("codigo", ""),
            p.get("nombre", ""),
            p.get("unidades_total", ""),
            ultima.get("fecha", ""),
            ultima.get("proximo_contacto") or "",
        ],
    )
    nombre_cliente = (
        "".join(c for c in (d.get("cliente") or "cliente") if c.isalnum() or c in " -_")[:40].strip() or "cliente"
    )
    return _descargar(libro, f"cliente-{nombre_cliente}-{date.today().isoformat()}.xlsx")


@bp.route("/api/exportar/resultados.xlsx")
def api_exportar_resultados():
    """Excel de la pestaña Resultados (?dias=30|90|365, sin dias = todo): hoja Resumen,
    hoja Clientes (un renglon por cliente+producto con su resultado mas reciente) y
    hoja Historial (todos los seguimientos)."""
    dias, valido = dias_de_la_peticion()
    if not valido:
        return jsonify({"error": "dias debe ser positivo"}), 400
    datos = datos_resultados(dias)
    historial = negocio.listar_seguimientos(dias)
    nombres = nombres_producto()

    libro = openpyxl.Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    periodo = f"Últimos {dias} días" if dias else "Todo el historial"
    resumen.append(["Resultados de los seguimientos", periodo, f"Generado el {date.today().isoformat()}"])
    resumen["A1"].font = openpyxl.styles.Font(bold=True, size=13)
    resumen.append([])
    resumen.append(["Casos atendidos (cliente + producto)", datos["total"]])
    resumen.append(["Clientes distintos", datos["clientes"]])
    tasa = datos["tasa_recuperacion"]
    resumen.append(
        ["Recuperación (volvieron a comprar de los casos resueltos)", f"{tasa}%" if tasa is not None else "—"]
    )
    resumen.append([])
    resumen.append(["Resultado más reciente de cada caso", "Casos"])
    _encabezado_en_negrita(resumen[resumen.max_row])
    for clave, etiqueta in ETIQUETAS_RESULTADO.items():
        resumen.append([etiqueta, datos["totales"][clave]])
    _anchos(resumen, [58, 18])

    def nota_del_caso(c):
        if c.get("por_factura"):
            return f"Factura nueva detectada (antes: {ETIQUETAS_RESULTADO.get(c['antes'], c['antes'])})"
        return c["nota"] or ""

    clientes = libro.create_sheet("Clientes")
    clientes.append(
        [
            "Cliente",
            "Código",
            "Producto",
            "Resultado",
            "Fecha",
            "Recordar el",
            "Intentos sin respuesta",
            "Nota / motivo",
        ]
    )
    _encabezado_en_negrita(clientes[1])
    for c in datos["casos"]:
        clientes.append(
            [
                c["clave_cliente"],
                c["codigo"],
                nombres.get(c["codigo"], ""),
                ETIQUETAS_RESULTADO.get(c["resultado"], c["resultado"]),
                c["fecha"][:10],
                c["proximo_recordatorio"] or "",
                c["intentos"] or "",
                nota_del_caso(c),
            ]
        )
    clientes.freeze_panes = "A2"
    clientes.auto_filter.ref = clientes.dimensions
    _anchos(clientes, [36, 12, 42, 20, 12, 13, 12, 46])

    detalle = libro.create_sheet("Historial")
    detalle.append(["Fecha", "Cliente", "Código", "Producto", "Resultado", "Recordar el", "Nota / motivo"])
    _encabezado_en_negrita(detalle[1])
    for d in historial:
        detalle.append(
            [
                d["fecha"][:16].replace("T", " "),
                d["clave_cliente"],
                d["codigo"],
                nombres.get(d["codigo"], ""),
                ETIQUETAS_RESULTADO.get(d["resultado"], d["resultado"]),
                d["proximo_recordatorio"] or "",
                d["nota"] or "",
            ]
        )
    detalle.freeze_panes = "A2"
    detalle.auto_filter.ref = detalle.dimensions
    _anchos(detalle, [17, 36, 12, 42, 20, 13, 46])

    return _descargar(libro, f"resultados-seguimientos-{date.today().isoformat()}.xlsx")
